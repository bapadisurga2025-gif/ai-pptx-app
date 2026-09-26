import streamlit as st
from google import genai
from pptx import Presentation
from pptx.util import Inches, Pt
import os
import pypdf
import pandas as pd

# Konfigurasi halaman
st.set_page_config(page_title="AI PPTX Executive Generator", layout="centered")

st.title("AI PPTX Generator - Template Cloner & Executive Style")
st.write("Sistem otomatis pengolahan data EB, Jaskug, dan Ritel dengan replikasi struktur template PowerPoint Anda.")

# Buat folder penyimpanan permanen jika belum ada
UPLOAD_DIR = "uploaded_data"
os.makedirs(UPLOAD_DIR, exist_ok=True)

def save_uploaded_file(uploaded_file, prefix):
    if uploaded_file is not None:
        file_path = os.path.join(UPLOAD_DIR, f"{prefix}_{uploaded_file.name}")
        with open(file_path, "wb") as f:
            f.write(uploaded_file.getbuffer())
        return file_path
    return None

def read_saved_file(file_path):
    if not file_path or not os.path.exists(file_path):
        return ""
    content = ""
    try:
        ext = file_path.split('.').[-1].lower()
        if ext == 'pdf':
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    content += text + "\n"
        elif ext in ['xlsx', 'xls']:
            df = pd.read_excel(file_path)
            content = df.to_string()
        elif ext == 'csv':
            df = pd.read_csv(file_path)
            content = df.to_string()
    except Exception as e:
        content = f"[Gagal membaca {file_path}: {e}]"
    return content

# Session state initialization
if "prompt" not in st.session_state:
    st.session_state.prompt = ""
if "output_path" not in st.session_state:
    st.session_state.output_path = None

prompt = st.text_area("Prompt / Instruksi Khusus EM", value=st.session_state.prompt, key="prompt_input")
st.session_state.prompt = prompt

st.markdown("---")
st.subheader("📁 Upload Data & Foto Pendukung (Admin / SPV)")

col1, col2 = st.columns(2)
with col1:
    f_eb = st.file_uploader("Upload File EB", type=["csv", "xlsx", "pdf"], key="eb_f")
    if f_eb: st.session_state.path_eb = save_uploaded_file(f_eb, "eb")
with col2:
    p_eb = st.file_uploader("Upload Foto EB", type=["png", "jpg", "jpeg"], key="eb_p")
    if p_eb: st.session_state.photo_eb = save_uploaded_file(p_eb, "photo_eb")

col3, col4 = st.columns(2)
with col3:
    f_jaskug = st.file_uploader("Upload File Jaskug", type=["csv", "xlsx", "pdf"], key="jaskug_f")
    if f_jaskug: st.session_state.path_jaskug = save_uploaded_file(f_jaskug, "jaskug")
with col4:
    p_jaskug = st.file_uploader("Upload Foto Jaskug", type=["png", "jpg", "jpeg"], key="jaskug_p")
    if p_jaskug: st.session_state.photo_jaskug = save_uploaded_file(p_jaskug, "photo_jaskug")

col5, col6 = st.columns(2)
with col5:
    f_ritel = st.file_uploader("Upload File Ritel", type=["csv", "xlsx", "pdf"], key="ritel_f")
    if f_ritel: st.session_state.path_ritel = save_uploaded_file(f_ritel, "ritel")
with col6:
    p_ritel = st.file_uploader("Upload Foto Ritel", type=["png", "jpg", "jpeg"], key="ritel_p")
    if p_ritel: st.session_state.photo_ritel = save_uploaded_file(p_ritel, "photo_ritel")

st.markdown("---")
uploaded_template = st.file_uploader("Upload Template PowerPoint (.pptx)", type=["pptx"], key="template_file")
if uploaded_template:
    st.session_state.path_template = save_uploaded_file(uploaded_template, "template")

st.markdown("---")
if st.button("Generate Executive PPTX"):
    if not prompt:
        st.warning("Mohon masukkan instruksi terlebih dahulu!")
    else:
        with st.spinner("Menyusun presentasi sesuai template dan data dengan Gemini AI..."):
            try:
                data_eb = read_saved_file(st.session_state.get("path_eb"))
                data_jaskug = read_saved_file(st.session_state.get("path_jaskug"))
                data_ritel = read_saved_file(st.session_state.get("path_ritel"))

                combined_data = f"""
                DATA EB: {data_eb[:1500]}
                DATA JASKUG: {data_jaskug[:1500]}
                DATA RITEL: {data_ritel[:1500]}
                """

                client = genai.Client(api_key=st.secrets["GEMINI_API_KEY"])
                response = client.models.generate_content(
                    model='gemini-2.5-flash',
                    contents=f"""Anda adalah konsultan manajemen senior. Berdasarkan instruksi '{prompt}' dan data berikut:
                    {combined_data}
                    
                    Buatkan materi presentasi terstruktur yang terbagi ke dalam beberapa slide. 
                    Setiap slide harus mengikuti format persis ini agar dapat diparsing sistem:
                    ---SLIDE---
                    KATEGORI: [Judul Panel Utama, misal: Kinerja KANAL EB]
                    METRIK: [Ringkasan Angka kunci, misal: 731.97 Miliar | 88.1%]
                    POIN_1: [Poin pertama ringkas]
                    POIN_2: [Poin kedua ringkas]
                    POIN_3: [Poin ketiga ringkas]
                    """
                )
                ai_output = response.text

                # Inisialisasi Presentation menggunakan Template yang diunggah
                if "path_template" in st.session_state and os.path.exists(st.session_state.path_template):
                    prs = Presentation(st.session_state.path_template)
                    base_layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
                else:
                    prs = Presentation()
                    prs.slide_width = Inches(13.333)
                    prs.slide_height = Inches(7.5)
                    base_layout = prs.slide_layouts[6]

                slides_data = ai_output.split("---SLIDE---")
                generated_count = 0

                for s_data in slides_data:
                    if "KATEGORI:" in s_data:
                        try:
                            lines = s_data.strip().split("\n")
                            cat, met, p1, p2, p3 = "Executive Overview", "", "", "", ""
                            for line in lines:
                                if "KATEGORI:" in line: cat = line.replace("KATEGORI:", "").strip()
                                elif "METRIK:" in line: met = line.replace("METRIK:", "").strip()
                                elif "POIN_1:" in line: p1 = line.replace("POIN_1:", "").strip()
                                elif "POIN_2:" in line: p2 = line.replace("POIN_2:", "").strip()
                                elif "POIN_3:" in line: p3 = line.replace("POIN_3:", "").strip()

                            if generated_count < len(prs.slides):
                                slide = prs.slides[generated_count]
                            else:
                                slide = prs.slides.add_slide(base_layout)
                            
                            generated_count += 1

                            filled_placeholders = 0
                            for shape in slide.shapes:
                                if shape.has_text_frame:
                                    if filled_placeholders == 0:
                                        shape.text_frame.text = cat
                                        filled_placeholders += 1
                                    elif filled_placeholders == 1:
                                        body_text = f"METRIK UTAMA: {met}\n\n• {p1}\n• {p2}\n• {p3}"
                                        shape.text_frame.text = body_text
                                        filled_placeholders += 1

                            if filled_placeholders < 2:
                                txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.0))
                                tf = txBox.text_frame
                                tf.word_wrap = True
                                p = tf.paragraphs[0]
                                p.text = f"Kategori: {cat}"
                                p.font.bold = True
                                p.font.size = Pt(20)
                                
                                p2_elem = tf.add_paragraph()
                                p2_elem.text = f"Metrik Utama: {met}"
                                p2_elem.font.size = Pt(16)
                                
                                for p_text in [p1, p2, p3]:
                                    if p_text:
                                        pt_elem = tf.add_paragraph()
                                        pt_elem.text = f"• {p_text}"
                                        pt_elem.font.size = Pt(14)

                        except Exception as ex:
                            continue

                st.session_state.output_path = "output_template_cloned.pptx"
                prs.save(st.session_state.output_path)
                st.success("Berhasil! Presentasi berhasil disusun dengan mereplikasi gaya dan template yang Anda unggah.")
            except Exception as e:
                st.error(f"Terjadi kesalahan: {e}")

if st.session_state.output_path and os.path.exists(st.session_state.output_path):
    with open(st.session_state.output_path, "rb") as f:
        st.download_button(
            label="Download File PPTX Sesuai Template",
            data=f,
            file_name="Presentasi_Sesuai_Template.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
